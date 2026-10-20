"""Write output/electricians.csv, output/electricians.xlsx and output/dashboard.html."""
import csv, hashlib, json, re, shutil
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from common import BASE, TODAY, db

OUT = BASE / "output"
COLS = [
    ("id", "ID"), ("priority", "Priority"), ("name", "Company"), ("owner", "Owner/contact (Impressum)"),
    ("legal_form", "Legal form"), ("street", "Street"), ("postcode", "Postcode"), ("town", "Town"),
    ("region", "Region"), ("dist_gersthofen", "Km from Gersthofen"), ("dist_center", "Km from region centre"),
    ("phone_landline", "Phone (landline)"), ("phone_mobile", "Phone (mobile)"), ("phones_all", "All phones"),
    ("email", "Email"), ("emails_all", "All emails"), ("website", "Website"), ("contact_form", "Contact form URL"),
    ("impressum_url", "Impressum URL"), ("size", "Size"), ("employees", "Employees (if stated)"),
    ("meister", "Meisterbetrieb"), ("knx_status", "KNX status"), ("knx_specialist", "KNX/smart-home specialist"),
    ("knx_snippet", "KNX mention (excerpt)"), ("services", "Other services"), ("brands", "Brands mentioned"),
    ("gs_rating", "Gelbe Seiten rating"), ("gs_reviews", "Gelbe Seiten reviews"),
    ("google_rating", "Google rating"), ("google_reviews", "Google reviews"),
    ("data_quality", "Data quality"), ("contact_method", "Recommended channel"), ("contact_reason", "Reason"),
    ("best_time", "Best time to call"), ("opener", "Opener (DE)"), ("phone_src", "Phone source"),
    ("email_src", "Email source"), ("source_urls", "Directory/OSM sources"), ("website_status", "Website check"),
    ("geo_approx", "Location approximate"), ("collected", "First collected"), ("updated", "Updated"),
]


def flat(v):
    if isinstance(v, list):
        return "; ".join(str(x) for x in v)
    if isinstance(v, bool):
        return "yes" if v else "no"
    return "" if v is None else v


def load():
    con = db()
    leads = [json.loads(r[0]) for r in con.execute("SELECT data FROM lead")]
    for l in leads:
        l["source_urls"] = [s["url"] for s in l["sources"]]
    leads.sort(key=lambda l: (-l["priority"], l["dist_gersthofen"]))
    return leads


PWA_SRC = BASE / "pwa"


def build_apps(data):
    """output/pwa/ (installable app + leads.json), output/dashboard.html (single file, data embedded)
    and output/knx-partner-leads.html (claude.ai page) -- all from the sources in pwa/."""
    payload = {"date": TODAY, "leads": data}
    dest = OUT / "pwa"
    if dest.exists():
        shutil.rmtree(dest)
    shutil.copytree(PWA_SRC, dest)
    leads_json = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    (dest / "leads.json").write_text(leads_json, encoding="utf-8")
    h = hashlib.sha256()
    for f in sorted(dest.rglob("*")):
        if f.is_file() and f.name != "sw.js":
            h.update(f.read_bytes())
    sw = dest / "sw.js"
    sw.write_text(sw.read_text(encoding="utf-8").replace("__VERSION__", h.hexdigest()[:12]), encoding="utf-8")

    page = (PWA_SRC / "index.html").read_text(encoding="utf-8")
    page = re.sub(r"(?s)<!--pwa-->.*?<!--/pwa-->\n?", "", page)
    css = (PWA_SRC / "styles.css").read_text(encoding="utf-8")
    js = (PWA_SRC / "app.js").read_text(encoding="utf-8")
    embed = ("window.EMBEDDED_DATE=" + json.dumps(TODAY) + ";window.EMBEDDED_LEADS="
             + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";").replace("</", "<\\/")
    page = page.replace('<link rel="stylesheet" href="styles.css">', "<style>\n" + css + "</style>")
    page = page.replace('<script src="app.js"></script>', "<script>" + embed + "</script>\n<script>\n" + js + "</script>")
    (OUT / "dashboard.html").write_text(page, encoding="utf-8")
    # claude.ai page: the platform supplies the document skeleton
    art = re.sub(r"(?is)<!doctype html>\s*<html[^>]*>\s*<head>\s*<meta charset[^>]*>\s*<meta name=\"viewport\"[^>]*>\s*", "", page)
    art = re.sub(r"(?is)</head>\s*<body>", "", art)
    art = re.sub(r"(?is)</body>\s*</html>\s*$", "", art)
    (OUT / "knx-partner-leads.html").write_text(art, encoding="utf-8")


def main():
    OUT.mkdir(exist_ok=True)
    leads = load()
    with open(OUT / "electricians.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([h for _, h in COLS])
        for l in leads:
            w.writerow([flat(l.get(k)) for k, _ in COLS])

    wb = Workbook()
    ws = wb.active
    ws.title = "Electricians"
    ws.append([h for _, h in COLS])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for l in leads:
        ws.append([flat(l.get(k)) for k, _ in COLS])
    widths = {"name": 38, "opener": 80, "contact_reason": 50, "website": 32, "email": 30, "knx_snippet": 50}
    for i, (k, _) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(k, 16)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    ex = wb.create_sheet("Excluded")
    ex.append(["ID", "Name", "Town", "Reason"])
    for r in db().execute("SELECT data FROM excluded"):
        x = json.loads(r[0])
        ex.append([x["id"], x["name"], x["town"], x["reason"]])
    wb.save(OUT / "electricians.xlsx")

    keep = ["id", "priority", "name", "owner", "legal_form", "street", "postcode", "town", "region", "dist_gersthofen",
            "phone_landline", "phone_mobile", "email", "website", "contact_form", "size", "meister", "knx_status",
            "knx_specialist", "knx_snippet", "services", "brands", "gs_rating", "gs_reviews", "data_quality",
            "contact_method", "contact_reason", "best_time", "opener", "source_urls", "impressum_url"]
    data = [{k: l.get(k) for k in keep} for l in leads]
    build_apps(data)
    print("exported", len(leads), "leads ->", OUT)


if __name__ == "__main__":
    main()
